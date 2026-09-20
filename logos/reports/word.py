"""Render a Report as a Word document (python-docx)."""
import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .model import Report

STATUS_COLORS = {  # (background, text)
    "OK": ("CDEFE0", "146B45"),
    "MISMATCH": ("FBD5E3", "9B1C4A"),
    "NEEDS_REVIEW": ("FFEDB5", "6B4E00"),
}
ROW_TINT = {"Mismatch": "FBD5E3", "Missing": "FFEDB5"}
HEADER_FILL = "F3EEFC"
BRAND = RGBColor(0xD9, 0x70, 0x3A)
MUTED = RGBColor(0x5B, 0x4F, 0x73)


def _shade(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell(cell, text, width_cm, fill=None):
    cell.text = str(text)
    cell.width = Cm(width_cm)
    if fill:
        _shade(cell, fill)


def _grid(doc, headers, rows, widths_cm, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text, width in zip(table.rows[0].cells, headers, widths_cm):
        _set_cell(cell, text, width, HEADER_FILL)
        cell.paragraphs[0].runs[0].bold = True
    for values, fill in zip(rows, row_fills or [None] * len(rows)):
        for cell, text, width in zip(table.add_row().cells, values, widths_cm):
            _set_cell(cell, text, width, fill)
    return table


def _add_header(doc, report):
    tagline = doc.add_paragraph()
    brand = tagline.add_run("Logos")
    brand.bold = True
    brand.font.color.rgb = BRAND
    tagline.add_run(" - SI vs draft BL comparison report").font.color.rgb = MUTED
    doc.add_heading(report.title, level=1)

    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    for label, value in report.meta:
        label_cell, value_cell = meta.add_row().cells
        _set_cell(label_cell, label, 5, HEADER_FILL)
        _set_cell(value_cell, value, 12)


def _add_verdict(doc, report):
    background, text = STATUS_COLORS[report.status]
    doc.add_paragraph()
    cell = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    _set_cell(cell, report.verdict, 17, background)
    run = cell.paragraphs[0].runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(text)


def _add_escalations(doc, report):
    doc.add_heading("Open escalations", level=2)
    for escalation in report.escalations:
        doc.add_paragraph(escalation.message)
        if escalation.evidence:
            evidence = doc.add_paragraph().add_run(escalation.evidence)
            evidence.font.name = "Courier New"
            evidence.font.size = Pt(8)


def render_word(report: Report) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)

    _add_header(doc, report)
    _add_verdict(doc, report)

    doc.add_heading("Field comparison", level=2)
    _grid(doc, ["Field", "SI (source of truth)", "Draft BL", "Result"],
          [(r.label, r.si, r.bl, r.result) for r in report.rows], [4, 5.5, 5.5, 2],
          [ROW_TINT.get(r.result) for r in report.rows])

    if report.escalations:
        _add_escalations(doc, report)

    doc.add_heading("Edit history", level=2)
    if report.history:
        _grid(doc, ["Field", "Old", "New", "Editor", "When", "Reason"],
              [(e.label, e.old, e.new, e.editor, e.timestamp, e.reason) for e in report.history],
              [3, 2.6, 2.6, 2.2, 2.6, 4])
    else:
        doc.add_paragraph("No edits recorded for this case")

    footer = doc.add_paragraph()
    note = footer.add_run(f"Generated {report.generated_at} by Logos. Case {report.email_id}.")
    note.font.size = Pt(8)
    note.font.color.rgb = MUTED

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
