"""Turn an attachment (txt/pdf/docx/xlsx) into plain text."""
import io

MIN_READABLE_CHARS = 40


def attachment_text(inbox, path: str) -> str:
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "txt":
        return inbox.read_text(path)
    data = inbox.read_bytes(path)
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    raise ValueError(f"unsupported attachment type: {ext}")


def is_readable(text: str) -> bool:
    return len(text.strip()) >= MIN_READABLE_CHARS
