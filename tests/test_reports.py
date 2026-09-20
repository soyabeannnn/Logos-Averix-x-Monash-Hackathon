import io
import sys
from pathlib import Path

import pytest
import sqlalchemy as sa
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from logos import api, db, reports, service  # noqa: E402
from logos.reports import build_report  # noqa: E402

GOOD = {
    "shipper": "ACME & SONS <LTD>", "consignee": "BETA CO", "notify_party": "GAMMA",
    "port_of_loading": "PORT KLANG", "port_of_discharge": "CALLAO",
    "container_count": 1, "gross_weight_kg": 21577.0,
}


def make_detail(**overrides):
    rows = [{"field": f, "si": v, "bl": v, "match": True} for f, v in GOOD.items()]
    rows[1] = {"field": "consignee", "si": "BETA CO", "bl": "OTHER CO", "match": False}
    rows[2] = {"field": "notify_party", "si": "GAMMA", "bl": None, "match": None}
    detail = {
        "id": "email_900", "subject": "Check docs & confirm", "sender": "a@b.c", "category": "BL_COMPARISON",
        "confidence": 0.95, "status": "MISMATCH", "processed_at": "2026-09-19T08:51:42+00:00",
        "shipment_ref": "5ALT-01226", "resolved_by": None, "resolved_at": None,
        "comparison": rows, "reasons": [], "edit_log": [],
    }
    return {**detail, **overrides}


# ---- report content (pure) ----
def test_report_formats_values_and_flags_rows():
    report = build_report(make_detail(), generated_at="now")
    by_label = {r.label: r for r in report.rows}
    assert by_label["Gross Weight (kg)"].si == "21,577"
    assert by_label["Notify Party"].bl == "missing" and by_label["Notify Party"].result == "Missing"
    assert by_label["Consignee"].result == "Mismatch"
    assert dict(report.meta)["Classification confidence"] == "95%"


def test_report_verdict_text_follows_status_and_resolution():
    assert "1 of 7 fields differ" in build_report(make_detail()).verdict
    resolved = build_report(make_detail(status="OK", resolved_by="Ana", resolved_at="2026-09-19T09:00:00+00:00"))
    assert resolved.verdict.startswith("Matched.") and "Resolved by Ana" in resolved.verdict


def test_report_history_includes_field_and_category_edits():
    log = [
        {"field": "category", "doc": "EMAIL", "old_value": "SPAM", "new_value": "BL_COMPARISON",
         "editor": "Ana", "reason": "wrong", "timestamp": "2026-09-19T09:00:00+00:00"},
        {"field": "consignee", "doc": "BL", "old_value": "OTHER CO", "new_value": "BETA CO",
         "editor": "Ana", "reason": "typo", "timestamp": "2026-09-19T09:05:00+00:00"},
    ]
    history = build_report(make_detail(edit_log=log)).history
    assert history[0].label == "Category" and history[0].old == "Spam" and history[0].new == "Comparison"
    assert history[1].label == "Consignee (BL)" and history[1].new == "BETA CO"


# ---- renderers ----
def pdf_text(data):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)


def test_pdf_renders_content_and_survives_markup_characters():
    detail = make_detail(reasons=[{"message": "BL: Notify Party missing", "evidence": "line <1> & more"}])
    data = reports.FORMATS["pdf"].render(build_report(detail))
    text = pdf_text(data)
    assert data.startswith(b"%PDF")
    assert "Check docs & confirm" in text and "ACME & SONS <LTD>" in text
    assert "OTHER CO" in text and "Mismatch" in text
    assert "No edits recorded for this case" in text and "BL: Notify Party missing" in text


def test_word_renders_tables_and_empty_history_message():
    data = reports.FORMATS["docx"].render(build_report(make_detail()))
    doc = Document(io.BytesIO(data))
    cells = {c.text for t in doc.tables for row in t.rows for c in row.cells}
    assert {"Consignee", "BETA CO", "21,577", "Mismatch", "Missing", "a@b.c"} <= cells
    assert any("No edits recorded for this case" in p.text for p in doc.paragraphs)


# ---- endpoint ----
@pytest.fixture()
def client():
    db.set_engine(sa.create_engine("sqlite://", poolclass=sa.pool.StaticPool,
                                   connect_args={"check_same_thread": False}))
    base = {
        "sender": "a@b.c", "subject": "s", "body": "b", "attachments": [], "confidence": 0.9,
        "rationale": "", "shipment_ref": None, "processed_at": "2026-09-19T08:00:00+00:00",
        "si_fields": GOOD, "bl_fields": {**GOOD, "consignee": "OTHER"}, "si_evidence": {}, "bl_evidence": {},
        "si_text": "x", "bl_text": "y", "reasons": [], "resolved_by": None, "resolved_at": None,
        "escalated_by": None,
    }
    service.save_row({**base, "id": "cmp", "category": "BL_COMPARISON", "status": "MISMATCH"})
    service.save_row({**base, "id": "spam", "category": "SPAM", "status": "CLASSIFIED"})
    return TestClient(api.app)


@pytest.mark.parametrize("fmt,mime,ext", [
    ("pdf", "application/pdf", "pdf"),
    ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
])
def test_report_endpoint_serves_a_download(client, fmt, mime, ext):
    res = client.get(f"/emails/cmp/report?format={fmt}")
    assert res.status_code == 200 and res.headers["content-type"] == mime
    assert res.headers["content-disposition"] == f'attachment; filename="logos-cmp.{ext}"'
    assert len(res.content) > 500


def test_report_endpoint_rejects_bad_requests(client):
    assert client.get("/emails/cmp/report?format=exe").status_code == 400
    assert client.get("/emails/spam/report?format=pdf").status_code == 400
    assert client.get("/emails/missing/report?format=pdf").status_code == 404


def test_health_endpoint(client):
    assert client.get("/health").json() == {"status": "ok"}
